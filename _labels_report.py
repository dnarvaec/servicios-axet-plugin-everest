import openpyxl
from collections import Counter, defaultdict
p = r"c:\Users\jherrodr\Downloads\EverestAutomatizacion\demo-servicios-axet-plugin\casos de prueba\casos everest.xlsx"
w = openpyxl.load_workbook(p, read_only=True, data_only=True)
s = w.active
headers = [s.cell(1, c).value for c in range(1, s.max_column + 1)]
idx = {str(v).strip().lower(): i for i, v in enumerate(headers) if v is not None}
label_i = idx["labels"]
issue_i = idx.get("issue key")
summary_i = idx.get("summary")
groups = defaultdict(list)
for row in s.iter_rows(min_row=2, values_only=True):
    labels = str(row[label_i] or "").strip()
    low = labels.lower()
    if any(x in low for x in ("biometr", "huella", "dactilar")):
        kind = "Biometrico"
    elif "accesib" in low or "accessibility" in low:
        kind = "Accesibilidad"
    elif "funcional" in low or "functional" in low:
        kind = "Funcional"
    else:
        kind = "Sin clasificar"
    groups[kind].append((row[issue_i] if issue_i is not None else "", row[summary_i] if summary_i is not None else "", labels))
print("TOTAL", sum(len(v) for v in groups.values()))
for kind in ("Funcional", "Accesibilidad", "Biometrico", "Sin clasificar"):
    values = groups.get(kind, [])
    print(kind, len(values))
    for issue, summary, labels in values:
        print(issue, "|", labels, "|", summary)
