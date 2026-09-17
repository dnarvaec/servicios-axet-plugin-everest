import openpyxl

path = r"c:\Users\jherrodr\Downloads\EverestAutomatizacion\demo-servicios-axet-plugin\tests\automatizacion api\serenity rest\src\test\resources\datadriven\datadriven.xlsx"
wb = openpyxl.load_workbook(path)

common = {
    "expected.httpStatusCode": 502,
    "expected.statusCode": "502",
    "expected.statusDesc": "Error en el servicio del banco",
}

def add(sheet, rows):
    ws = wb[sheet]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    for row in rows:
        values = {**common, **row}
        values["Caso"] = ws.max_row
        ws.append([values.get(h, "") for h in headers])

add("avv_consulta_cliente", [
    {"operacion": "CONSULTA_CLIENTE", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "Cliente no encontrado en AVV"},
    {"operacion": "CONSULTA_CLIENTE", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 200, "expected.statusCode": "206", "expected.statusDesc": "Respuesta parcial"},
    {"operacion": "CONSULTA_CLIENTE", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 504, "expected.statusCode": "504", "expected.statusDesc": "Timeout del servicio AVV"},
])
add("avv_consulta_productos", [
    {"operacion": "CONSULTA_PRODUCTOS", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "El cliente no tiene productos en AVV"},
    {"operacion": "CONSULTA_PRODUCTOS", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 200, "expected.statusCode": "206", "expected.statusDesc": "Respuesta parcial"},
    {"operacion": "CONSULTA_PRODUCTOS", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 502, "expected.statusCode": "502", "expected.statusDesc": "Error de conectividad con AVV"},
])
add("bdb_consulta_general", [
    {"operacion": "CONSULTA_GENERAL", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "El cliente no tiene productos en BdB"},
    {"operacion": "CONSULTA_GENERAL", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 200, "expected.statusCode": "206", "expected.statusDesc": "Respuesta parcial"},
    {"operacion": "CONSULTA_GENERAL", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 502, "expected.statusCode": "502", "expected.statusDesc": "Error en el servicio del banco"},
])
add("avv_cartera_detallada", [
    {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "26456554", "expected.httpStatusCode": 200, "expected.statusCode": "200", "expected.statusDesc": "Obligacion en mora"},
    {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "26456554", "expected.httpStatusCode": 502, "expected.statusCode": "502", "expected.statusDesc": "Error de conectividad con AVV"},
    {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "26456554", "expected.httpStatusCode": 504, "expected.statusCode": "504", "expected.statusDesc": "Timeout del servicio AVV"},
])
add("bdb_cartera_detallada", [
    {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "559536650", "expected.httpStatusCode": 422, "expected.statusCode": "422", "expected.statusDesc": "Error de negocio en BdB"},
    {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "559536650", "expected.httpStatusCode": 504, "expected.statusCode": "504", "expected.statusDesc": "Timeout del servicio BdB"},
    {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "559536650", "expected.httpStatusCode": 200, "expected.statusCode": "200", "expected.statusDesc": "Inactive"},
])
add("avv_tc_detallada", [
    {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 200, "expected.statusCode": "200", "expected.statusDesc": "Tarjeta bloqueada"},
    {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 502, "expected.statusCode": "502", "expected.statusDesc": "Error de conectividad con AVV"},
    {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 504, "expected.statusCode": "504", "expected.statusDesc": "Timeout del servicio AVV"},
])
add("bdb_tc_detallada", [
    {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 200, "expected.statusCode": "200", "expected.statusDesc": "CANCELADA"},
    {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 422, "expected.statusCode": "422", "expected.statusDesc": "Error de negocio en BdB"},
    {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 502, "expected.statusCode": "502", "expected.statusDesc": "Error en el servicio del banco"},
])
add("avv_cdt_detallado", [
    {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "26456554", "expected.httpStatusCode": 200, "expected.statusCode": "200", "expected.statusDesc": "CDT vigente"},
    {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "26456554", "expected.httpStatusCode": 502, "expected.statusCode": "502", "expected.statusDesc": "Error de conectividad con AVV"},
    {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "26456554", "expected.httpStatusCode": 504, "expected.statusCode": "504", "expected.statusDesc": "Timeout del servicio AVV"},
])
add("bdb_cdt_detallado", [
    {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "0000000026456554", "expected.httpStatusCode": 422, "expected.statusCode": "422", "expected.statusDesc": "Error de negocio en BdB"},
    {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "0000000026456554", "expected.httpStatusCode": 504, "expected.statusCode": "504", "expected.statusDesc": "Timeout del servicio BdB"},
    {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "0000000026456554", "expected.httpStatusCode": 200, "expected.statusCode": "200", "expected.statusDesc": "VIGENTE"},
])

wb.save(path)
print("remaining functional cases appended")
