from __future__ import annotations

from copy import copy
from pathlib import Path
import re
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
CASES_EXCEL = ROOT / "casos de prueba" / "casos everest.xlsx"
PROJECT = ROOT / "tests" / "automatizacion api" / "serenity rest"
DATADRIVEN = PROJECT / "src" / "test" / "resources" / "datadriven" / "datadriven.xlsx"
FEATURE = PROJECT / "src" / "test" / "resources" / "features" / "oficinas" / "oficinas-actualizaciones.feature"


HEADERS = [
    "Caso",
    "header.Content-Type",
    "header.Authorization",
    "header.X-Trace-Id",
    "header.X-Origin-Bank",
    "header.X-Destination-Bank",
    "operacion",
    "obj_operacion.tipoDocumento",
    "obj_operacion.numeroDocumento",
    "obj_operacion.numDocumento",
    "obj_operacion.numeroTarjeta",
    "obj_operacion.causal",
    "obj_operacion.datosActualizar.nombre1",
    "obj_operacion.datosActualizar.apellido1",
    "obj_operacion.datosActualizar.celular",
    "obj_operacion.datosActualizar.correoElectronico",
    "obj_operacion.datosActualizar.ingresosMensuales",
    "expected.httpStatusCode",
    "expected.statusCode",
    "expected.statusDesc",
]

BANK_BY_HU = {
    "AVV": ("BAVV", "avv"),
    "BDB": ("BBOG", "bdb"),
    "BPO": ("BPOP", "bpop"),
    "OCC": ("BOCC", "occ"),
}

CARD_BY_BANK = {
    "BAVV": "4532123456789012",
    "BBOG": "4551234567890123",
    "BOCC": "5307100045685945",
    "BPOP": "5307100045685945",
}


def normalize(text: str) -> str:
    value = unicodedata.normalize("NFD", text.lower())
    return "".join(ch for ch in value if unicodedata.category(ch) != "Mn")


def slug(text: str) -> str:
    value = normalize(text)
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-")[:180]


def case_title(summary: str, story: str) -> str:
    prefix = f"[{story}]"
    if summary.startswith(prefix):
        return summary[len(prefix):].strip()
    return summary.strip()


def headers_of(worksheet):
    return {str(cell.value).strip(): index for index, cell in enumerate(worksheet[1]) if cell.value}


def cell_value(row, headers, name: str) -> str:
    index = headers.get(name)
    return str(row[index].value or "").strip() if index is not None else ""


def bank_for_story(story: str):
    for token, bank_info in BANK_BY_HU.items():
        if f"-{token}" in story:
            return bank_info
    raise ValueError(f"No se pudo identificar banco para {story}")


def operation_for_story(story: str):
    if story.startswith("HU-201"):
        return "BLOQUEO_TD_DEFINITIVO", "bloqueo_td"
    if story.startswith("HU-203"):
        return "ACTUALIZACION_DATOS", "actualizacion_datos"
    raise ValueError(f"HU no soportada: {story}")


def status_for_case(story: str, bank: str, operation: str, text: str):
    if "no encontr" in text or "inexistente" in text:
        return 404, "Tarjeta no encontrada" if operation == "BLOQUEO_TD_DEFINITIVO" else "Cliente no encontrado"
    if "ya bloqueada" in text or "bloqueo activo" in text or "conflicto" in text:
        return 409, "Tarjeta ya bloqueada"
    if "operacion no soportada" in text or "operacion no esperada" in text or "dato obligatorio" in text or "header" in text or "datos invalidos" in text:
        return 400, "Solicitud invalida"
    if "timeout" in text or "falla" in text or "error interno" in text or "credencial" in text or "certificado" in text or "ip no autorizada" in text or "resiliencia" in text or "comunicacion" in text:
        return 502, "Error tecnico del banco"
    if operation == "BLOQUEO_TD_DEFINITIVO" and bank == "BAVV":
        return 503, "Servicio no disponible para esta entidad"
    if operation == "ACTUALIZACION_DATOS" and bank != "BAVV":
        return 503, "Servicio no disponible para esta entidad"
    return 200, "OK"


def trace_for_status(status: int) -> str:
    if status == 400:
        return "trace-missing-destination"
    if status == 502:
        return "trace-technical-error"
    return "trace-001"


def document_for_status(status: int, operation: str, bank: str) -> str:
    if status == 404:
        return "00000000"
    if status == 409:
        return "11111111"
    if status == 502:
        return "99999999"
    if operation == "ACTUALIZACION_DATOS" and bank == "BAVV":
        return "86068761"
    return "12345678"


def build_row(case_number: int, story: str, bank: str, operation: str, title: str, status: int, status_desc: str):
    document = document_for_status(status, operation, bank)
    data = {header: "" for header in HEADERS}
    data.update({
        "Caso": case_number,
        "header.Content-Type": "application/json",
        "header.Authorization": "Bearer test-token",
        "header.X-Trace-Id": trace_for_status(status),
        "header.X-Origin-Bank": "BBOG",
        "header.X-Destination-Bank": "" if status == 400 else bank,
        "operacion": operation,
        "obj_operacion.tipoDocumento": "CC",
        "expected.httpStatusCode": status,
        "expected.statusCode": str(status),
        "expected.statusDesc": status_desc,
    })
    if operation == "BLOQUEO_TD_DEFINITIVO":
        data["obj_operacion.numeroDocumento"] = document
        data["obj_operacion.numeroTarjeta"] = "1111111111111111" if status == 409 else CARD_BY_BANK[bank]
        data["obj_operacion.causal"] = "BL03" if "fraude" in normalize(title) else "BL01"
    else:
        if bank == "BOCC":
            data["obj_operacion.numDocumento"] = document
        else:
            data["obj_operacion.numeroDocumento"] = document
        data["obj_operacion.datosActualizar.nombre1"] = "JUAN"
        data["obj_operacion.datosActualizar.apellido1"] = "GARCIA"
        data["obj_operacion.datosActualizar.celular"] = "3001234567"
        data["obj_operacion.datosActualizar.correoElectronico"] = "juan.garcia@test.com"
        if status == 200:
            data["obj_operacion.datosActualizar.ingresosMensuales"] = 5000000
    return data


def load_automation_cases():
    workbook = openpyxl.load_workbook(CASES_EXCEL, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = headers_of(worksheet)
    cases = []
    for row in worksheet.iter_rows(min_row=2):
        story = cell_value(row, headers, "Story Linkages")
        test_type = cell_value(row, headers, "TestCase Type")
        if not (story.startswith("HU-201") or story.startswith("HU-203")):
            continue
        if test_type != "Funcional":
            continue
        summary = cell_value(row, headers, "Summary")
        description = cell_value(row, headers, "Description")
        expected = cell_value(row, headers, "Expected Result")
        searchable = normalize(" ".join([summary, description, expected]))
        if "biometr" in searchable or "biometric" in searchable:
            continue
        issue_key = cell_value(row, headers, "Issue Key")
        bank, prefix = bank_for_story(story)
        operation, suffix = operation_for_story(story)
        sheet_name = f"{prefix}_{suffix}"
        title = case_title(summary, story)
        status, status_desc = status_for_case(story, bank, operation, searchable)
        cases.append({
            "issue_key": issue_key,
            "story": story,
            "bank": bank,
            "operation": operation,
            "sheet_name": sheet_name,
            "title": title,
            "status": status,
            "status_desc": status_desc,
        })
    return cases


def build_sheet_rows(cases):
    sheets = {}
    for case in cases:
        rows = sheets.setdefault(case["sheet_name"], [])
        case_number = len(rows) + 1
        case["case_number"] = case_number
        rows.append(build_row(
            case_number,
            case["story"],
            case["bank"],
            case["operation"],
            case["title"],
            case["status"],
            case["status_desc"],
        ))
    return sheets


def copy_sheet_style(source, target):
    for column in range(1, len(HEADERS) + 1):
        cell = target.cell(row=1, column=column)
        source_cell = source.cell(row=1, column=min(column, source.max_column))
        if source_cell.has_style:
            cell.font = copy(source_cell.font)
            cell.fill = copy(source_cell.fill)
            cell.border = copy(source_cell.border)
            cell.alignment = copy(source_cell.alignment)
        target.column_dimensions[get_column_letter(column)].width = max(14, min(42, len(HEADERS[column - 1]) + 2))


def write_datadriven(sheets):
    workbook = openpyxl.load_workbook(DATADRIVEN)
    style_source = workbook[workbook.sheetnames[0]]
    for sheet_name, rows in sheets.items():
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(HEADERS)
        copy_sheet_style(style_source, worksheet)
        for item in rows:
            worksheet.append([item.get(header, "") for header in HEADERS])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, max_col=worksheet.max_column):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(
                    left=Side(style="thin", color="D3D3D3"),
                    right=Side(style="thin", color="D3D3D3"),
                    top=Side(style="thin", color="D3D3D3"),
                    bottom=Side(style="thin", color="D3D3D3"),
                )
    workbook.save(DATADRIVEN)


def write_feature(cases):
    lines = [
        "# language: es",
        "@oficinas @actualizaciones",
        "Característica: Oficinas - Bloqueo TD y actualizacion de datos",
        "",
    ]
    for case in cases:
        scenario = f"{case['story']} - {case['title']} - {case['issue_key']} - codigo {case['status']} - valida respuesta"
        lines.extend([
            f"  @{slug(case['issue_key'] + '-' + scenario)}",
            f"  Esquema del escenario: {scenario}",
            f"    Cuando ejecuta la operación nueva \"{case['operation']}\" del banco \"{case['bank']}\" en el caso <Caso>",
            "    Entonces la respuesta de la operación nueva coincide con lo esperado",
            "",
            "    Ejemplos:",
            f"      ##@externaldata@src/test/resources/datadriven/datadriven.xlsx@{case['sheet_name']}",
            "      | Caso |",
            f"      |{case['case_number']}|",
            "",
        ])
    FEATURE.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    automation_cases = load_automation_cases()
    sheet_rows = build_sheet_rows(automation_cases)
    write_datadriven(sheet_rows)
    write_feature(automation_cases)
    print(f"Casos funcionales no biometricos automatizados: {len(automation_cases)}")
    for sheet_name in sorted(sheet_rows):
        print(f"{sheet_name}: {len(sheet_rows[sheet_name])}")
    print(f"Feature: {FEATURE}")