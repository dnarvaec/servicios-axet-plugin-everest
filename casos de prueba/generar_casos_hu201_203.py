from __future__ import annotations

import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT / "casos de prueba" / "casos everest.xlsx"
TEMPLATE_PATH = ROOT / "insumos_cp" / "Formato Jira.xlsx"
SHEET_NAME = "Casos funcionales"
TODAY = "2026-09-23"


HU_201 = {
    "AVV": {
        "hu": "HU-201-ADP-AVV",
        "bank": "AV Villas",
        "component": "Adaptador AVV - Bloqueo TD Definitivo",
        "dest": "BAVV",
        "service": "servicio de bloqueo de medios de AV Villas con indicador de bloqueo definitivo",
        "success": "el bloqueo definitivo se confirma y se retorna el resultado normalizado al orquestador",
        "constant": "indBloqueo igual a S",
        "pending": "contrato AVV pendiente de confirmacion para campos de request, response y causales",
        "special": [
            ("validar el registro de errores de conectividad o servicio", "el servicio bancario informa un error", "el error queda registrado en Elastic"),
        ],
        "errors": [
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
        ],
    },
    "BDB": {
        "hu": "HU-201-ADP-BDB",
        "bank": "Banco de Bogota",
        "component": "Adaptador BdB - Bloqueo TD Definitivo",
        "dest": "BBOG",
        "service": "servicio de gestion de tarjeta debito de Banco de Bogota",
        "success": "la tarjeta queda con bloqueo definitivo y la respuesta se normaliza hacia el orquestador",
        "constant": "estado destino de tarjeta igual a C para bloqueo definitivo",
        "pending": "confirmacion de tarjeta completa o enmascarada y mapeo definitivo de causales",
        "special": [
            ("validar el contexto obligatorio de la operacion", "la solicitud contiene contexto de cliente, canal, terminal y trazabilidad", "el banco recibe el contexto requerido para procesar el bloqueo"),
            ("validar que las credenciales dependan del ambiente", "la integracion usa la credencial configurada para el ambiente activo", "la credencial no queda hardcodeada en la aplicacion"),
            ("validar el registro de errores del banco", "el banco retorna un error de conectividad o de negocio", "el error queda registrado en Elastic"),
        ],
        "errors": [
            ("confirmar la respuesta cuando la tarjeta no existe", "tarjeta no encontrada en el banco", "404"),
            ("confirmar la respuesta cuando la tarjeta ya esta bloqueada", "tarjeta con bloqueo activo previo", "409"),
            ("confirmar el mapeo de error de negocio", "error funcional retornado por el banco", "422"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar credenciales o acceso de integracion invalido", "credencial de integracion invalida o no autorizada", "502"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
    "BPO": {
        "hu": "HU-201-ADP-BPO",
        "bank": "Banco Popular",
        "component": "Adaptador BPO - Bloqueo TD Definitivo",
        "dest": "BPOP",
        "service": "servicio de bloqueo de tarjeta debito de Banco Popular",
        "success": "los prerequisitos de identidad y consulta de tarjeta se cumplen, el bloqueo se confirma y la respuesta se normaliza",
        "constant": None,
        "pending": "WSDL, endpoint de consulta de tarjeta, whitelist y mapeo causal pendientes de confirmacion",
        "prereq": True,
        "causal_success": True,
        "special": [
            ("detener el bloqueo cuando falla la verificacion de identidad", "el prerequisito de verificacion de identidad no es exitoso", "no se consulta la tarjeta ni se ejecuta el bloqueo"),
            ("detener el bloqueo cuando falla la consulta de tarjeta", "la verificacion de identidad es exitosa pero no se obtiene el identificador interno de la tarjeta", "no se ejecuta el bloqueo"),
            ("validar la seguridad requerida por el banco", "la operacion usa la politica de seguridad, certificado y contexto configurados para el ambiente", "el banco acepta el contexto de seguridad sin exponer credenciales en el codigo"),
        ],
        "errors": [
            ("confirmar la respuesta cuando la tarjeta no existe", "tarjeta no encontrada en el banco", "404"),
            ("confirmar el mapeo de error de negocio", "error de negocio retornado por el banco", "422"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar una IP no autorizada por el banco", "IP de salida no incluida en whitelist", "502"),
            ("manejar un error de certificado o seguridad", "certificado TLS o politica de seguridad rechazada", "502"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
    "OCC": {
        "hu": "HU-201-ADP-OCC",
        "bank": "Banco de Occidente",
        "component": "Adaptador OCC - Bloqueo TD Definitivo",
        "dest": "BOCC",
        "service": "servicio de bloqueo y cancelacion de tarjeta debito de Banco de Occidente",
        "success": "el banco confirma la cancelacion definitiva y la respuesta se normaliza con estado bloqueada",
        "constant": "accion de cancelacion definitiva para causales de robo, perdida o fraude",
        "pending": "valores exactos de BloqueoTipo, ProductoEstado y formato de tarjeta pendientes de confirmacion",
        "causal_success": True,
        "special": [
            ("validar el contexto requerido para solicitar el bloqueo", "la operacion contiene contexto de aplicacion, terminal, sesion, peticion, usuario y fecha", "el banco recibe el contexto completo documentado en la historia"),
            ("validar que el identificador de peticion sea unico", "se ejecutan dos solicitudes de bloqueo", "cada solicitud usa un identificador de peticion diferente"),
            ("validar que el usuario ejecutor provenga del contexto autorizado", "el asesor esta autenticado", "el usuario enviado al banco corresponde al contexto del asesor y no al cuerpo de la solicitud"),
        ],
        "errors": [
            ("confirmar la respuesta cuando la tarjeta no existe", "tarjeta no encontrada en el banco", "404"),
            ("confirmar la respuesta cuando el estado no permite bloqueo", "estado de tarjeta incompatible con bloqueo definitivo", "422"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar una falla de comunicacion con el banco", "servicio bancario no disponible", "502"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
}


HU_203 = {
    "AVV": {
        "hu": "HU-203-ADP-AVV",
        "bank": "AV Villas",
        "component": "Adaptador AVV - Actualizacion de Datos",
        "dest": "BAVV",
        "service": "servicio de actualizacion de datos de cliente de AV Villas",
        "success": "la actualizacion se confirma con codigo exitoso y mensaje normalizado",
        "fields": "celular, correo electronico, direccion, ciudad e informacion financiera enviada por el orquestador",
        "special": [
            ("construir el contexto tecnico obligatorio", "el contexto de autenticacion, funcionario, oficina, dispositivo, fecha y hora esta disponible", "el banco recibe todos los datos obligatorios documentados en la historia"),
            ("mapear tipos de documento permitidos", "tipos CC, CE, NIT y TI", "tipo de documento traducido al codigo requerido por el banco"),
            ("enviar solo campos presentes para actualizar", "solicitud parcial con celular y correo", "el banco recibe solo los datos informados, sin nulos ni vacios"),
            ("obtener los datos del asesor desde el contexto", "funcionario, oficina y tokens de sesion disponibles", "los datos del asesor se toman del contexto autorizado y no del cuerpo de la solicitud"),
        ],
        "errors": [
            ("confirmar el mapeo de error de negocio", "respuesta bancaria no exitosa", "422"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar una falla de comunicacion con el banco", "servicio bancario no disponible", "502"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
    "BDB": {
        "hu": "HU-203-ADP-BDB",
        "bank": "Banco de Bogota",
        "component": "Adaptador BdB - Actualizacion de Datos",
        "dest": "BBOG",
        "service": "servicio de actualizacion de informacion basica de Banco de Bogota",
        "success": "la actualizacion se confirma como exitosa y se retorna mensaje normalizado",
        "fields": "celular, correo electronico, telefono de residencia y telefono de oficina",
        "special": [
            ("construir el contexto obligatorio de la solicitud", "estan disponibles los datos de cliente, canal, empresa, red, terminal y trazabilidad", "el banco recibe el contexto obligatorio documentado en la historia"),
            ("generar un identificador unico por solicitud", "solicitud valida de actualizacion", "cada intento usa un identificador nuevo para trazabilidad"),
            ("enviar solo campos presentes para actualizar", "solicitud parcial con datos de contacto", "el cuerpo contiene solo los datos informados por el orquestador"),
            ("mapear correctamente los telefonos por tipo", "celular, residencia y oficina", "cada telefono queda clasificado segun su tipo funcional"),
        ],
        "errors": [
            ("confirmar el rechazo por datos invalidos", "datos de cliente invalidos", "400"),
            ("confirmar el rechazo por cliente no encontrado", "cliente inexistente en el banco", "400"),
            ("manejar credenciales o acceso de integracion invalido", "credencial o acceso no autorizado", "502"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
    "BPO": {
        "hu": "HU-203-ADP-BPO",
        "bank": "Banco Popular",
        "component": "Adaptador BPO - Actualizacion de Datos",
        "dest": "BPOP",
        "service": "servicio de actualizacion de datos de Banco Popular",
        "success": "la actualizacion se confirma en el MDM y se retorna resultado normalizado",
        "fields": "informacion de contacto, preferencias de comunicacion, informacion financiera y laboral habilitada",
        "special": [
            ("construir el sobre de integracion documentado", "estan disponibles el contexto del canal, banco, operador, cliente y transaccion", "el banco recibe el sobre con todos los datos obligatorios documentados en la historia"),
            ("construir el contexto del operador y del cliente", "JWT del asesor y datos del cliente disponibles", "el sobre de integracion separa correctamente operador y cliente"),
            ("generar identificadores unicos de transaccion", "solicitud valida de actualizacion", "cada intento queda trazable con identificadores nuevos"),
            ("enviar solo nodos con datos presentes", "solicitud parcial con datos de contacto", "el banco recibe solo estructuras correspondientes a los campos informados"),
        ],
        "errors": [
            ("confirmar el rechazo por datos invalidos", "datos de cliente invalidos", "400"),
            ("confirmar el rechazo por cliente no encontrado", "cliente inexistente en MDM", "400"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar un error interno del banco", "servicio bancario retorna error interno", "502"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
    "OCC": {
        "hu": "HU-203-ADP-OCC",
        "bank": "Banco de Occidente",
        "component": "Adaptador OCC - Actualizacion de Datos",
        "dest": "BOCC",
        "service": "servicio de actualizacion de datos de Banco de Occidente",
        "success": "el banco confirma la actualizacion y el adaptador retorna respuesta exitosa normalizada",
        "fields": "celular y correo electronico segun catalogo de contacto autorizado",
        "special": [
            ("construir el contexto requerido por el banco", "estan disponibles el token, aplicacion, usuario, sesion, peticion, terminal y fecha", "el banco recibe el contexto obligatorio documentado en la historia"),
            ("obtener token de integracion antes de actualizar datos", "credenciales de integracion vigentes", "la operacion usa un token valido obtenido por el adaptador"),
            ("reutilizar token vigente sin solicitar uno nuevo", "token cacheado aun vigente", "el adaptador usa el token en cache antes de su expiracion"),
            ("renovar token expirado y reintentar una vez", "token rechazado por expiracion", "el adaptador renueva el token y completa la actualizacion si el segundo intento es exitoso"),
        ],
        "errors": [
            ("confirmar el rechazo por datos invalidos", "datos de cliente invalidos", "400"),
            ("manejar error al obtener token de integracion", "no se puede autenticar con el banco", "502"),
            ("manejar token invalido despues del reintento", "segundo intento rechazado por autenticacion", "502"),
            ("rechazar una operacion no soportada", "operacion no esperada por el adaptador", "400"),
            ("manejar un timeout del banco", "servicio bancario sin respuesta dentro del tiempo configurado", "504"),
        ],
    },
}


def functional_case(hu, bank, action, condition, expected, component, priority="High"):
    return {
        "TestCase Type": "Funcional",
        "Summary": f"[{hu}] Como analista de pruebas quiero {action} para confirmar que {expected}",
        "Description": f"Validar desde una perspectiva funcional que, para {bank}, cuando {condition}, {expected}.",
        "Precondition": f"Ambiente de pruebas disponible para {bank}; orquestador y adaptador configurados; datos de negocio preparados.",
        "Priority": priority,
        "Step Summary": "1. Preparar datos de negocio y contexto autorizado. | 2. Ejecutar el caso de uso desde el orquestador. | 3. Verificar respuesta, estado funcional y trazabilidad esperada.",
        "Test Data": f"Banco: {bank}; condicion: {condition}.",
        "Expected Result": expected.capitalize() + ".",
        "Components": component,
    }


def accessibility_case(hu, bank, action, expected, component, focus):
    return {
        "TestCase Type": "Accesibilidad",
        "Summary": f"[{hu}] Como analista de pruebas quiero {action} para confirmar que {expected}",
        "Description": f"Validar que la informacion generada por el flujo de {bank} sea comprensible, consistente y usable por personas que dependen de mensajes claros o ayudas de lectura.",
        "Precondition": f"Resultado funcional del flujo de {bank} disponible para revision de contenido y presentacion.",
        "Priority": "Medium",
        "Step Summary": "1. Obtener una respuesta exitosa o de error representativa. | 2. Revisar claridad, orden, rotulado y ausencia de ambiguedad. | 3. Confirmar que el mensaje permite tomar accion sin conocimiento tecnico.",
        "Test Data": f"Foco de accesibilidad: {focus}.",
        "Expected Result": expected.capitalize() + ".",
        "Components": component,
    }


def biometric_cases(hu, bank, operation, component):
    return [
        functional_case(
            hu,
            bank,
            f"validar biometricos OK antes de {operation}",
            "la huella del cliente coincide y la autenticacion biometrica es aprobada",
            f"el sistema autoriza {operation} y continua el flujo de negocio",
            component,
            "High",
        ),
        functional_case(
            hu,
            bank,
            f"validar biometricos No OK antes de {operation}",
            "la huella del cliente no coincide o es rechazada por biometria",
            f"el sistema bloquea {operation} y retorna un rechazo de autenticacion biometrica",
            component,
            "High",
        ),
    ]


def bloqueo_cases(bank_data):
    hu = bank_data["hu"]
    bank = bank_data["bank"]
    component = bank_data["component"]
    operation = "el bloqueo definitivo de tarjeta debito"
    cases = [
        functional_case(hu, bank, f"ejecutar {operation} exitosamente", "la solicitud contiene cliente, tarjeta y causal validos", bank_data["success"], component),
        functional_case(hu, bank, "validar el enrutamiento por banco y operacion", f"la operacion corresponde al banco destino {bank_data['dest']}", "el adaptador selecciona el servicio bancario correcto sin procesar operaciones inesperadas", component, "Medium"),
        functional_case(hu, bank, "validar auditoria y trazabilidad de la operacion", "la operacion finaliza con exito o error", "se registra el request recibido y la respuesta o error antes de retornar al orquestador", component, "Medium"),
        functional_case(hu, bank, "validar enmascaramiento en registros y respuesta", "la operacion contiene numero de tarjeta", "la informacion sensible se protege en salidas de consulta humana y registros de auditoria", component, "High"),
        functional_case(hu, bank, "validar activacion de resiliencia ante fallas repetidas", "el banco presenta fallas consecutivas", "el mecanismo de proteccion evita llamadas indefinidas y retorna un error controlado", component, "Medium"),
    ]
    if bank_data.get("constant"):
        cases.append(functional_case(hu, bank, "validar la constante funcional de bloqueo definitivo", bank_data["constant"], "el bloqueo enviado al banco corresponde a un bloqueo definitivo", component, "High"))
    if bank_data.get("causal_success"):
        cases.append(functional_case(hu, bank, "ejecutar el bloqueo con las causales documentadas", "se ejecuta la operacion con BL01, BL02 y BL03 usando los datos exitosos de la coleccion", "el banco confirma el bloqueo para cada causal documentada en la coleccion", component, "High"))
    if bank_data.get("prereq"):
        cases.append(functional_case(hu, bank, "validar prerequisitos de identidad y consulta de tarjeta", "los pasos previos retornan respuesta exitosa", "el bloqueo se ejecuta solo despues de cumplir los prerequisitos", component, "High"))
    cases.extend(
        functional_case(hu, bank, action, condition, expected, component, "High")
        for action, condition, expected in bank_data.get("special", [])
    )
    cases.extend(
        functional_case(hu, bank, action, condition, f"se retorna codigo {code} al orquestador con mensaje funcional claro", component, "Medium")
        for action, condition, code in bank_data["errors"]
    )
    cases.extend(biometric_cases(hu, bank, operation, component))
    cases.extend([
        accessibility_case(hu, bank, "revisar la claridad del mensaje de bloqueo exitoso", "la confirmacion del bloqueo sea comprensible y no ambigua", component, "mensaje exitoso de bloqueo"),
        accessibility_case(hu, bank, "revisar la claridad de los mensajes de error de bloqueo", "el usuario entienda el error documentado para la operacion y la accion que debe tomar", component, "mensajes de error documentados en la historia"),
        accessibility_case(hu, bank, "revisar la legibilidad de tarjeta causal y estado", "los datos sensibles se presenten de forma protegida y entendible", component, "tarjeta enmascarada, causal y estado"),
    ])
    return cases


def actualizacion_cases(bank_data):
    hu = bank_data["hu"]
    bank = bank_data["bank"]
    component = bank_data["component"]
    operation = "la actualizacion de datos del cliente"
    cases = [
        functional_case(hu, bank, f"ejecutar {operation} exitosamente", f"la solicitud contiene datos actualizables validos: {bank_data['fields']}", bank_data["success"], component),
        functional_case(hu, bank, "validar el enrutamiento por banco y operacion", f"la operacion corresponde al banco destino {bank_data['dest']}", "el adaptador selecciona el servicio bancario correcto sin procesar operaciones inesperadas", component, "Medium"),
        functional_case(hu, bank, "validar auditoria y trazabilidad de la actualizacion", "la operacion finaliza con exito o error", "se registra el request recibido y la respuesta o error antes de retornar al orquestador", component, "Medium"),
        functional_case(hu, bank, "validar proteccion de credenciales y configuracion", "la operacion requiere credenciales, tokens o variables de ambiente", "no se usan credenciales hardcodeadas y la configuracion depende del ambiente activo", component, "High"),
        functional_case(hu, bank, "validar activacion de resiliencia ante fallas repetidas", "el banco presenta fallas consecutivas", "el mecanismo de proteccion evita llamadas indefinidas y retorna un error controlado", component, "Medium"),
    ]
    cases.extend(
        functional_case(hu, bank, action, condition, expected, component, "High")
        for action, condition, expected in bank_data["special"]
    )
    cases.extend(
        functional_case(hu, bank, action, condition, f"se retorna codigo {code} al orquestador con mensaje funcional claro", component, "Medium")
        for action, condition, code in bank_data["errors"]
    )
    cases.extend(biometric_cases(hu, bank, operation, component))
    cases.extend([
        accessibility_case(hu, bank, "revisar la claridad de la confirmacion de actualizacion", "la confirmacion indique que datos fueron actualizados sin lenguaje tecnico innecesario", component, "mensaje exitoso de actualizacion"),
        accessibility_case(hu, bank, "revisar la claridad de errores de datos invalidos", "el usuario entienda que dato debe corregir o validar", component, "mensajes de error funcional"),
        accessibility_case(hu, bank, "revisar la legibilidad de datos de contacto y preferencia", "telefonos, correos y preferencias se presenten con rotulos consistentes", component, "datos de contacto actualizados"),
    ])
    return cases


def all_cases():
    generated = []
    for data in HU_201.values():
        generated.extend((data["hu"], data["component"], case) for case in bloqueo_cases(data))
    for data in HU_203.values():
        generated.extend((data["hu"], data["component"], case) for case in actualizacion_cases(data))
    selected = []
    for hu_id, component, case in generated:
        summary = case["Summary"].lower()
        if case["TestCase Type"] == "Accesibilidad" or "biometric" in summary:
            selected.append((hu_id, component, case))
            continue

        text = case["Summary"] + " " + case["Expected Result"]
        has_explicit_status = re.search(
            r"(?i)(?:codigo\s+|retorna\s+(?:codigo\s+)?|HTTP\s+)(200|400|404|409|422|500|502|503|504)",
            text,
        )
        is_success = re.search(
            r"(?i)ejecutar .* exitosamente|respuesta exitosa|actualizacion exitosa|confirma el bloqueo|confirma la actualizacion|operacion exitosa",
            text,
        )
        if not has_explicit_status and not is_success:
            continue
        if is_success and not has_explicit_status:
            case["Expected Result"] = case["Expected Result"].rstrip(".") + ". Se retorna codigo 200."
        selected.append((hu_id, component, case))
    return selected


def header_map(ws):
    return {str(ws.cell(row=1, column=col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(row=1, column=col).value}


def next_issue_number(ws, headers):
    issue_col = headers.get("Issue Key")
    highest = 0
    if not issue_col:
        return 1
    for row in range(2, ws.max_row + 1):
        value = str(ws.cell(row=row, column=issue_col).value or "")
        match = re.fullmatch(r"EV-(\d+)", value.strip())
        if match:
            highest = max(highest, int(match.group(1)))
    return highest + 1


def copy_row_style(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        source = ws.cell(row=source_row, column=col)
        target = ws.cell(row=target_row, column=col)
        if source.has_style:
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)


def default_style(ws, row):
    border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def ensure_workbook():
    if OUTPUT_PATH.exists():
        return openpyxl.load_workbook(OUTPUT_PATH)
    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    workbook.active.title = SHEET_NAME
    return workbook


def append_cases():
    wb = ensure_workbook()
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    headers = header_map(ws)
    summary_col = headers["Summary"]
    existing_summaries = {
        str(ws.cell(row=row, column=summary_col).value or "").strip()
        for row in range(2, ws.max_row + 1)
    }

    issue_number = next_issue_number(ws, headers)
    start_row = ws.max_row + 1
    appended = 0
    by_type = {"Funcional": 0, "Accesibilidad": 0}
    by_hu = {}

    for hu_id, component, case in all_cases():
        if case["Summary"] in existing_summaries:
            continue
        row = ws.max_row + 1
        if row > 2:
            copy_row_style(ws, row - 1, row)
        else:
            default_style(ws, row)

        values = {
            "Issue Key": f"EV-{issue_number}",
            "Summary": case["Summary"],
            "Description": case["Description"],
            "Precondition": case["Precondition"],
            "Status": "To Do",
            "Priority": case.get("Priority", "Medium"),
            "Assignee": "",
            "Reporter": "",
            "Estimated Time": "",
            "Labels": f"everest,{case['TestCase Type'].lower()},oficinas,{hu_id.lower()}",
            "Components": component,
            "Sprint": "Por definir",
            "Fix Versions": "v1.0.0",
            "Is Shareable Step": "No",
            "Shareable Testcase Issue Key": "",
            "Shareable Testcase Version No.": "",
            "Step Summary": case["Step Summary"],
            "Test Data": case["Test Data"],
            "Expected Result": case["Expected Result"],
            "Version": "1",
            "Folders": hu_id,
            "TestCase Type": case["TestCase Type"],
            "Created By": "QA Automation Team",
            "Created Date": TODAY,
            "Updated By": "QA Automation Team",
            "Updated Date": TODAY,
            "Story Linkages": hu_id,
            "Comment Count": "0",
            "Attachment Count": "0",
            "Story Count": "1",
        }
        for name, col in headers.items():
            ws.cell(row=row, column=col, value=values.get(name, ""))
            ws.cell(row=row, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        existing_summaries.add(case["Summary"])
        issue_number += 1
        appended += 1
        by_type[case["TestCase Type"]] += 1
        by_hu[hu_id] = by_hu.get(hu_id, 0) + 1

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(col)
        if ws.column_dimensions[letter].width is None or ws.column_dimensions[letter].width < 12:
            ws.column_dimensions[letter].width = 20

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUTPUT_PATH)
    return appended, by_type, by_hu, start_row, ws.max_row


if __name__ == "__main__":
    appended, by_type, by_hu, start_row, end_row = append_cases()
    print(f"Casos agregados: {appended}")
    print(f"Rango agregado: filas {start_row}-{end_row}" if appended else "Sin filas nuevas")
    print(f"Funcionales: {by_type['Funcional']}")
    print(f"Accesibilidad: {by_type['Accesibilidad']}")
    for hu_id, count in sorted(by_hu.items()):
        print(f"{hu_id}: {count}")