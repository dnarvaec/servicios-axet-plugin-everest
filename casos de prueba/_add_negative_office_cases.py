import openpyxl
from copy import copy

path = r"c:\Users\jherrodr\Downloads\EverestAutomatizacion\demo-servicios-axet-plugin\tests\automatizacion api\serenity rest\src\test\resources\datadriven\datadriven.xlsx"
wb = openpyxl.load_workbook(path)

cases = {
    "avv_consulta_cliente": [
        {"operacion": "CONSULTA_CLIENTE", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "9999999999", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "Cliente no encontrado en AVV"},
        {"operacion": "OPERACION_INVALIDA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_CLIENTE", "obj_operacion.tipoDocumento": "CC", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "avv_consulta_productos": [
        {"operacion": "CONSULTA_PRODUCTOS", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "00000000", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "Cliente no encontrado en AVV"},
        {"operacion": "CONSULTA_INVALIDA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "86068761", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_PRODUCTOS", "obj_operacion": "", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "avv_cartera_detallada": [
        {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "9999999999", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "Obligacion no encontrada en AVV"},
        {"operacion": "CARTERA_DESCONOCIDA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "26456554", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "OBL-ABC-123", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "avv_tc_detallada": [
        {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4999999999999999", "expected.httpStatusCode": 404, "expected.statusCode": "404", "expected.statusDesc": "Tarjeta no encontrada en AVV"},
        {"operacion": "TC_INVALIDA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "123456", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "avv_cdt_detallado": [
        {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "CDT-999999", "expected.httpStatusCode": 404, "expected.statusCode": "404", "expected.statusDesc": "CDT no encontrado en AVV"},
        {"operacion": "CDT_NO_VALIDO", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "26456554", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "bdb_consulta_general": [
        {"operacion": "CONSULTA_GENERAL", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "9876543210", "expected.httpStatusCode": 206, "expected.statusCode": "206", "expected.statusDesc": "Cliente no encontrado en BdB"},
        {"operacion": "CONSULTA_ERRONEA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_GENERAL", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "bdb_cartera_detallada": [
        {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "000000000", "expected.httpStatusCode": 404, "expected.statusCode": "404", "expected.statusDesc": "Obligacion no encontrada en BdB"},
        {"operacion": "OPERACION_NO_VALIDA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroObligacion": "559536650", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_DETALLADA_CARTERA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "bdb_tc_detallada": [
        {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "0000000000000000", "expected.httpStatusCode": 404, "expected.statusCode": "404", "expected.statusDesc": "Tarjeta no encontrada en BdB"},
        {"operacion": "OP_TC_ERRONEA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.referenciaTarjeta": "4111111111111111", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_DETALLADA_TC", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
    "bdb_cdt_detallado": [
        {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "CDT-000000000", "expected.httpStatusCode": 404, "expected.statusCode": "404", "expected.statusDesc": "CDT no encontrado en BdB"},
        {"operacion": "OP_CDT_DESCONOCIDA", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "obj_operacion.numeroProducto": "0000000026456554", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
        {"operacion": "CONSULTA_DETALLADA_CDT", "obj_operacion.tipoDocumento": "CC", "obj_operacion.numeroDocumento": "12345678", "expected.httpStatusCode": 400, "expected.statusCode": "302", "expected.statusDesc": "Error de validacion de estructura"},
    ],
}

for sheet_name, negative_cases in cases.items():
    ws = wb[sheet_name]
    headers = {ws.cell(row=1, column=col).value: col for col in range(1, ws.max_column + 1)}
    next_case = max(int(ws.cell(row=row, column=headers["Caso"]).value) for row in range(2, ws.max_row + 1) if ws.cell(row=row, column=headers["Caso"]).value) + 1
    for case in negative_cases:
        case["Caso"] = next_case
        next_case += 1
        ws.append([case.get(header, "") for header in headers])

wb.save(path)
print("negative cases added")
